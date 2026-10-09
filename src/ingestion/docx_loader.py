"""DOCX document loader for ingesting structured reports into CanonicalEvidenceRecords.

Uses python-docx to parse incident reports, threat briefings, and SOC documentation,
preserving headings, paragraphs, and tabular structures.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from src.ingestion.evidence_schema import CanonicalEvidenceRecord, ModalityType

logger = logging.getLogger(__name__)


class DOCXLoader:
    """Loads and extracts text and structured tables from DOCX files."""

    def load_file(
        self,
        file_path: str | Path,
        source_dataset: str = "documents_docx",
    ) -> list[CanonicalEvidenceRecord]:
        """Load a DOCX file and return canonical evidence records for sections and tables.

        Args:
            file_path: Path to the .docx file.
            source_dataset: Dataset category name.

        Returns:
            List of CanonicalEvidenceRecord objects.
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        try:
            import docx
        except ImportError:
            raise ImportError("python-docx is required. Install with: pip install python-docx")

        doc = docx.Document(str(file_path))
        records: list[CanonicalEvidenceRecord] = []

        logger.info("Loading DOCX file: %s", file_path)

        current_heading = "Document Header"
        current_paragraph_buffer: list[str] = []
        chunk_idx = 0

        # Extract paragraphs grouped by headings
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Check if paragraph is a heading
            if para.style and para.style.name and para.style.name.startswith("Heading"):
                if current_paragraph_buffer:
                    # Flush previous section
                    content = f"[DOCX SECTION: {current_heading}]\n" + "\n\n".join(current_paragraph_buffer)
                    evidence_id = f"DOCX-{file_path.stem}-{chunk_idx + 1:03d}"
                    records.append(
                        CanonicalEvidenceRecord(
                            evidence_id=evidence_id,
                            source_file=file_path.name,
                            source_dataset=source_dataset,
                            modality=ModalityType.DOCX.value,
                            content=content,
                            raw_reference=str(file_path),
                            metadata={"heading": current_heading, "type": "section_text"},
                            chunk_id=chunk_idx,
                            event_type="document_section",
                        )
                    )
                    chunk_idx += 1
                    current_paragraph_buffer = []

                current_heading = text
            else:
                current_paragraph_buffer.append(text)

        # Flush trailing paragraphs
        if current_paragraph_buffer:
            content = f"[DOCX SECTION: {current_heading}]\n" + "\n\n".join(current_paragraph_buffer)
            evidence_id = f"DOCX-{file_path.stem}-{chunk_idx + 1:03d}"
            records.append(
                CanonicalEvidenceRecord(
                    evidence_id=evidence_id,
                    source_file=file_path.name,
                    source_dataset=source_dataset,
                    modality=ModalityType.DOCX.value,
                    content=content,
                    raw_reference=str(file_path),
                    metadata={"heading": current_heading, "type": "section_text"},
                    chunk_id=chunk_idx,
                    event_type="document_section",
                )
            )
            chunk_idx += 1

        # Extract tables as structured canonical evidence
        for tbl_idx, table in enumerate(doc.tables):
            table_rows: list[list[str]] = []
            for row in table.rows:
                table_rows.append([cell.text.strip() for cell in row.cells])

            if not table_rows:
                continue

            # Format table as Markdown text
            header = table_rows[0]
            header_line = "| " + " | ".join(header) + " |"
            separator = "| " + " | ".join(["---"] * len(header)) + " |"
            data_lines = [
                "| " + " | ".join(r) + " |" for r in table_rows[1:]
            ]

            table_text = f"[DOCX TABLE {tbl_idx + 1}: Source {file_path.name}]\n" + "\n".join(
                [header_line, separator] + data_lines
            )

            evidence_id = f"DOCX-TBL-{file_path.stem}-{tbl_idx + 1:03d}"
            records.append(
                CanonicalEvidenceRecord(
                    evidence_id=evidence_id,
                    source_file=file_path.name,
                    source_dataset=source_dataset,
                    modality=ModalityType.DOCX.value,
                    content=table_text,
                    raw_reference=str(file_path),
                    metadata={
                        "table_index": tbl_idx + 1,
                        "rows": len(table_rows),
                        "columns": len(header),
                        "type": "structured_table",
                    },
                    chunk_id=chunk_idx,
                    event_type="document_table",
                )
            )
            chunk_idx += 1

        logger.info(
            "Extracted %d canonical evidence records from DOCX: %s",
            len(records), file_path.name,
        )
        return records
